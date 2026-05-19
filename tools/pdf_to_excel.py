"""PDF -> Excel tool (Mistral OCR) as a Flask Blueprint."""
import os
import re
import io
from flask import Blueprint, request, jsonify, send_file, render_template
from mistralai import Mistral

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# Metadata used by the portal landing page to render the tool card.
TOOL = {
    "slug": "pdf-to-excel",
    "name": "PDF → Excel",
    "description": "Extract tables & text from any PDF into a structured Excel workbook using Mistral OCR.",
    "icon": "📄",
    "url": "/tools/pdf-to-excel",
}

bp = Blueprint(
    "pdf_to_excel",
    __name__,
    url_prefix="/tools/pdf-to-excel",
)

MISTRAL_API_KEY = os.environ.get("MISTRAL_API_KEY", "")


def parse_markdown_page(markdown_text):
    """Split a page's markdown into blocks: 'table' or 'text'."""
    blocks = []
    lines = markdown_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if re.match(r"^\s*\|", line):
            table_lines = []
            while i < len(lines) and re.match(r"^\s*\|", lines[i]):
                table_lines.append(lines[i])
                i += 1
            blocks.append({"type": "table", "content": table_lines})
        else:
            text_lines = []
            while i < len(lines) and not re.match(r"^\s*\|", lines[i]):
                text_lines.append(lines[i])
                i += 1
            text = "\n".join(text_lines).strip()
            if text:
                blocks.append({"type": "text", "content": text})
    return blocks


def parse_table_rows(table_lines):
    """Parse markdown table lines into list of row lists."""
    rows = []
    for line in table_lines:
        if re.match(r"^\s*\|[\s\-:]+\|", line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def write_page_to_sheet(ws, blocks):
    """Write parsed blocks to an openpyxl worksheet."""
    current_row = 1
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9E1F2")

    for block in blocks:
        if block["type"] == "table":
            rows = parse_table_rows(block["content"])
            for r_idx, row in enumerate(rows):
                for c_idx, cell_val in enumerate(row):
                    cell = ws.cell(row=current_row, column=c_idx + 1, value=cell_val)
                    if r_idx == 0:
                        cell.font = header_font
                        cell.fill = header_fill
                        cell.alignment = Alignment(horizontal="center")
                current_row += 1
            current_row += 1
        elif block["type"] == "text":
            for line in block["content"].split("\n"):
                line = line.strip()
                if line:
                    ws.cell(row=current_row, column=1, value=line)
                    current_row += 1
            current_row += 1


def auto_fit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)


@bp.route("/")
def index():
    return render_template("pdf_to_excel.html", tool=TOOL)


@bp.route("/convert", methods=["POST"])
def convert():
    if "pdf" not in request.files:
        return jsonify({"error": "No PDF file provided"}), 400

    pdf_file = request.files["pdf"]
    if not pdf_file.filename.lower().endswith(".pdf"):
        return jsonify({"error": "File must be a PDF"}), 400

    api_key = os.environ.get("MISTRAL_API_KEY", "")
    if not api_key:
        return jsonify({"error": "MISTRAL_API_KEY not configured on server"}), 500

    pdf_bytes = pdf_file.read()
    client = Mistral(api_key=api_key)

    uploaded = client.files.upload(
        file={"file_name": pdf_file.filename, "content": pdf_bytes},
        purpose="ocr",
    )
    signed_url = client.files.get_signed_url(file_id=uploaded.id)

    ocr_response = client.ocr.process(
        model="mistral-ocr-latest",
        document={"type": "document_url", "document_url": signed_url.url},
    )

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for page_idx, page in enumerate(ocr_response.pages):
        ws = wb.create_sheet(title=f"Page {page_idx + 1}")
        blocks = parse_markdown_page(page.markdown)
        write_page_to_sheet(ws, blocks)
        auto_fit_columns(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    base_name = os.path.splitext(pdf_file.filename)[0]
    excel_filename = f"{base_name}_ocr.xlsx"

    try:
        client.files.delete(file_id=uploaded.id)
    except Exception:
        pass

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=excel_filename,
    )
